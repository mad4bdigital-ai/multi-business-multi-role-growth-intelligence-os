#!/usr/bin/env python3
"""Offline, bounded PDF/image/DOCX structural audit. No OCR or source writes."""
from __future__ import annotations
import argparse
import json
import re
from pathlib import Path
try:
    import fitz
    from PIL import Image, ImageChops
    from docx import Document
except ImportError as exc:
    raise SystemExit("Install PyMuPDF, Pillow and python-docx") from exc

EXTENSIONS={".pdf":"pdf",".docx":"docx",".jpeg":"image",".jpg":"image",".png":"image",".webp":"image"}
DRAFT_RE=re.compile(r"(?:temp[-_ ]placeholder|will be changed later|coming soon|lorem ipsum|under construction)",re.I)

def pdf_status(path:Path,max_pages:int)->dict:
    details=[]
    try:
        with fitz.open(path) as doc:
            if doc.is_encrypted or len(doc)>max_pages:
                return dict(type="pdf",status="QUARANTINED",reason="ENCRYPTED_OR_PAGE_LIMIT")
            for n,page in enumerate(doc,1):
                try:
                    has_text=bool(page.get_text("text").strip())
                    pix=page.get_pixmap(matrix=fitz.Matrix(.32,.32),alpha=False,colorspace=fitz.csRGB)
                    im=Image.frombytes("RGB",(pix.width,pix.height),pix.samples)
                    visible=ImageChops.difference(im,Image.new("RGB",im.size,"white")).getbbox() is not None
                    details.append(dict(page=n,rendered=True,has_text_layer=has_text,has_visible_ink=visible))
                except Exception:
                    details.append(dict(page=n,rendered=False,has_text_layer=False,has_visible_ink=False))
    except Exception:
        return dict(type="pdf",status="QUARANTINED",reason="DOCUMENT_OPEN_FAILED")
    return dict(type="pdf",status="REVIEW_REQUIRED",pages=len(details),
                no_text_pages=sum(not x["has_text_layer"] for x in details),
                render_failures=sum(not x["rendered"] for x in details),
                possibly_blank_pages=sum(not x["has_visible_ink"] for x in details),
                page_checks=details,semantic_verified=False)

def image_status(path:Path)->dict:
    try:
        with Image.open(path) as im:
            im.load()
            w,h=im.size
            return dict(type="image",status="REVIEW_REQUIRED",decode_pass=True,
                        width=w,height=h,very_tall=h>3*w,visual_semantic_verified=False)
    except Exception:
        return dict(type="image",status="QUARANTINED",decode_pass=False)

def docx_status(path:Path)->dict:
    try:
        doc=Document(path)
        text="\n".join(p.text for p in doc.paragraphs)
        text+="\n"+"\n".join(" ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
        flags=[]
        if DRAFT_RE.search(text):flags.append("PLACEHOLDER_COPY")
        if len(text.strip())<500:flags.append("SHORT_REFERENCE")
        if doc.inline_shapes and len(text.strip())<100:flags.append("VISUAL_ONLY_OR_IMAGE_DOMINANT")
        return dict(type="docx",status="REVIEW_REQUIRED",paragraphs=len(doc.paragraphs),
                    tables=len(doc.tables),inline_images=len(doc.inline_shapes),text_length=len(text),
                    flags=flags,semantic_verified=False)
    except Exception:
        return dict(type="docx",status="QUARANTINED",reason="INVALID_DOCX")

def audit(root:Path,recursive:bool=False,max_files:int=500,max_mb:int=50,max_pages:int=500,include_names:bool=False)->dict:
    root=Path(root).resolve()
    if not root.is_dir():raise ValueError("INVALID_SOURCE_DIRECTORY")
    if min(max_files,max_mb,max_pages)<1:raise ValueError("INVALID_BUDGET")
    paths=sorted((p for p in (root.rglob("*") if recursive else root.iterdir())
                  if p.is_file() and p.suffix.lower() in EXTENSIONS),key=lambda p:str(p))
    if len(paths)>max_files:raise ValueError("FILE_COUNT_LIMIT")
    records=[]
    for i,p in enumerate(paths,1):
        kind=EXTENSIONS[p.suffix.lower()]
        if p.is_symlink() or not p.resolve().is_relative_to(root):
            rec=dict(type=kind,status="QUARANTINED",reason="SYMLINK_OR_PATH_ESCAPE")
        elif p.stat().st_size>max_mb*1024*1024:
            rec=dict(type=kind,status="QUARANTINED",reason="FILE_SIZE_LIMIT")
        else:
            rec=pdf_status(p,max_pages) if kind=="pdf" else image_status(p) if kind=="image" else docx_status(p)
        rec["asset_ref"]=f"local-{i:04d}"
        if include_names:rec["private_source_path"]=str(p.relative_to(root))
        records.append(rec)
    return dict(contract="mad4b.reference-source-audit.v1",status="STRUCTURAL_REVIEW_ONLY",
                source_data_copied=False,approval_granted=False,publication_authorized=False,
                source_revision_verified=False,semantic_visual_review_complete=False,
                counts={kind:sum(x["type"]==kind for x in records) for kind in ("pdf","docx","image")},
                pdf_pages=sum(x.get("pages",0) for x in records),
                pdf_pages_without_text=sum(x.get("no_text_pages",0) for x in records),
                pdf_render_failures=sum(x.get("render_failures",0) for x in records),
                records=records)

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("directory",type=Path)
    parser.add_argument("--recursive",action="store_true")
    parser.add_argument("--include-private-names",action="store_true")
    parser.add_argument("--max-files",type=int,default=500)
    parser.add_argument("--max-mb",type=int,default=50)
    parser.add_argument("--max-pages",type=int,default=500)
    parser.add_argument("--output",type=Path)
    args=parser.parse_args()
    result=audit(args.directory,args.recursive,args.max_files,args.max_mb,args.max_pages,args.include_private_names)
    encoded=json.dumps(result,ensure_ascii=False,indent=2)+"\n"
    if args.output:
        args.output.parent.mkdir(parents=True,exist_ok=True)
        args.output.write_text(encoded,encoding="utf-8")
        print("AUDIT_WRITTEN")
    else:print(encoded)
if __name__=="__main__":
    main()
